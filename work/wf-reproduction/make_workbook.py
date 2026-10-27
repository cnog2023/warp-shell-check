"""
make_workbook.py -- the numbers behind the figure pairs, as a workbook (Part 3, at Carlos's request).

One sheet per figure of the paper (4 to 10): at every radius of our grid line (0.1 to 29.5 m, 0.2 m
steps) the paper's curve (digitised from the PDF's vector paths, sampled every 2 mm, interpolated
linearly to our radii), ours (WarpFactory at the published live script's configuration, Part 1
export), their difference and ratio as live formulas, and, where the quantity follows from the 1D
construction alone, the replica rebuilt at the smoothing fitted to the paper's figures
(fit_smoothing.py: density span 6100, pressure span 3540). A Summary sheet lists every curve's
extrema and formulas for the RMS and maximum difference; a "Smoothing fit" sheet holds the fit; a
"Part 2 exact vs sampled" sheet holds the per-case energy-condition numbers of
work/matrix/run-fine/exact_cases.csv.

Inputs: paper_curves.csv (digitise_paper.py), the slices and profiles_f1.csv in work/matrix/run-fine,
comparison.csv, fit_smoothing.csv. Output: comparison.xlsx next to this script.

Usage: python make_workbook.py
"""
import csv
import sys
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(ROOT / "work" / "shell"))
from fit_smoothing import construct  # noqa: E402
from side_by_side import C_SI, R_S, load_case, load_profiles  # noqa: E402
import hoop_stress as hs  # noqa: E402
import replicate_profiles as rp  # noqa: E402

RUN = ROOT / "work" / "matrix" / "run-fine"
FONT = "Arial"
SCI = "0.000E+00"
FIX = "0.0000"
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")


def paper_curves():
    out = {}
    with open(HERE / "paper_curves.csv", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            out.setdefault((row["figure"], row["curve"]), []).append((float(row["r_m"]), float(row["value"])))
    for k, pts in out.items():
        pts.sort()
        out[k] = (np.array([p[0] for p in pts]), np.array([p[1] for p in pts]))
    return out


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_note(ws, row, text, ncols=8):
    ws.cell(row=row, column=1, value=text).font = Font(name=FONT, italic=True)
    ws.cell(row=row, column=1).fill = NOTE_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 48


def figure_sheet(wb, title, note, r, curves, fmt):
    """curves: list of (label, paper_values or None, ours_values, replica_values or None)."""
    ws = wb.create_sheet(title)
    write_note(ws, 1, note, ncols=max(8, 1 + 5 * len(curves)))
    ws.cell(row=2, column=1, value="r [m]")
    col = 2
    blocks = []
    for label, paper, ours, rep in curves:
        heads = [f"{label}: paper", f"{label}: ours", "difference (paper - ours)", "ratio paper / ours"]
        if rep is not None:
            heads.append(f"{label}: replica at the figures' smoothing")
        for h in heads:
            ws.cell(row=2, column=col + heads.index(h), value=h)
        blocks.append((col, label, paper is not None, rep is not None))
        n = len(r)
        for i in range(n):
            row = 3 + i
            cp, co, cd, cr = (get_column_letter(col + k) for k in range(4))
            if paper is not None:
                ws.cell(row=row, column=col, value=float(paper[i])).number_format = fmt
            ws.cell(row=row, column=col + 1, value=float(ours[i])).number_format = fmt
            if paper is not None:
                ws.cell(row=row, column=col + 2, value=f"={cp}{row}-{co}{row}").number_format = fmt
                ws.cell(row=row, column=col + 3, value=f'=IF(ABS({co}{row})<1E-300,"",{cp}{row}/{co}{row})').number_format = "0.0000"
            if rep is not None:
                ws.cell(row=row, column=col + 4, value=float(rep[i])).number_format = fmt
        col += len(heads)
    for i, ri in enumerate(r):
        ws.cell(row=3 + i, column=1, value=float(ri)).number_format = "0.0"
    style_header(ws, 2, col - 1)
    ws.row_dimensions[2].height = 45
    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 8
    for c in range(2, col):
        ws.column_dimensions[get_column_letter(c)].width = 16
    for row in ws.iter_rows(min_row=3, max_row=2 + len(r), max_col=col - 1):
        for cell in row:
            cell.font = Font(name=FONT)
    return ws, blocks, 3, 2 + len(r)


def main():
    pc = paper_curves()
    st, w, prof = load_case(RUN, 1), load_case(RUN, 5), load_profiles(RUN)
    r = st["r"]

    def P(fig, name):
        rr, vv = pc[(str(fig), name)]
        return np.interp(r, rr, vv)

    def prof_at(key):
        return np.interp(r, prof["r"], prof[key])

    # the replica at the smoothing fitted to the paper's figures (and the metric from it)
    base = rp.build(smoothFactor=4000)
    rep = construct(base, 6100, 3540)
    rr_, rho_r, Pr_r, Pt_r, _ = hs.identity_tangential(rep)
    with np.errstate(divide="ignore", invalid="ignore"):
        B_r = 1.0 / (1.0 - 2 * rp.G * rep["M_kg"] / (rep["r_m"] * rp.c ** 2))
    R = lambda arr, rgrid=rep["r_m"]: np.interp(r, rgrid, arr)
    rep_rho = R(rep["rhosmooth_kg_m3"] * C_SI ** 2); rep_P = R(rep["Psmooth_Pa"]); rep_g00 = R(rep["A_g00"]); rep_g22 = R(B_r)
    rep_Pt = np.interp(r, rr_, Pt_r); rep_Pr = np.interp(r, rr_, Pr_r)
    with np.errstate(divide="ignore"):
        g00_sch = -(1 - R_S / r); g22_sch = np.where(r > R_S, 1 / (1 - R_S / r), np.nan)

    wb = Workbook()
    ws = wb.active; ws.title = "Read me"
    lines = [
        "The numbers behind the figure pairs of work/wf-reproduction (Fuchs et al. 2024, arXiv:2405.02709, Figures 4 to 10, against ours).",
        "",
        "Paper columns: the paper's curves read off the PDF's vector paths (digitise_paper.py), each panel calibrated on its tick marks (baselines that must be zero read zero to 1e-4 of the axis range); the paper's lines are sampled every 2 mm and are interpolated linearly to our radii. Figure 8 is a raster image in the PDF and has no paper column.",
        "Ours columns: WarpFactory at the published live script's configuration (W1_Warp_Shell.mlx: R1 = 10 m, R2 = 20 m, M = 4.49e27 kg, smoothFactor 4000, 300 x 300 x 5 cells at 0.2 m), the Part 1 export, read on the grid column at x = -0.1 m (the grid has no x = 0 column); the two outermost cells are dropped. Static shell for Figures 4 to 7, shift 0.02 for Figures 8 to 10.",
        "Replica columns: the construction rebuilt in NumPy (work/shell/replicate_profiles.py, validated against WarpFactory to 1e-13) at the smoothing fitted to the paper's own smoothed inputs: density span 6100 samples, pressure span 3540 (ratio 1.72, the paper's footnote 7), where the live script uses 7160 and 4000 (ratio 1.79). Only quantities that follow from the 1D construction have this column: the smoothed inputs, the radial pressure, the tangential pressure by the equilibrium identity, g00 and g22.",
        "Difference and ratio columns are formulas; Excel computes them when the file opens. Units: J/m^3 for every stress-energy quantity (energy density times c^2, momentum density times c), metric components dimensionless, r in metres.",
        "",
        "What the numbers show: the paper's and our curves have the same shape at the same radii in every panel; the paper's amplitudes are 1 to 11 % higher, most where the feature is sharpest, and the replica at the fitted smoothing reproduces the paper's values (smoothed inputs to 0.02 %, hoop-stress spike to 0.14 %, wall width to 3 mm, g22 peak to 1e-4). The unsmoothed inputs agree to 1e-4 without any fitting: it is the same shell. The published live script is therefore not the exact configuration behind the figures (register row C10); the physics is unchanged. Write-up: work/wf-reproduction/RESULTS.md.",
        "",
        "Sheets: Summary (extrema per curve, RMS and maximum differences); Fig 4 to Fig 10; Smoothing fit; Part 2 exact vs sampled (work/matrix/run-fine/exact_cases.csv: warpax's exact margins against WarpFactory's sampled maps, Part 2).",
        "Built 27 October 2026 by work/wf-reproduction/make_workbook.py. Paper figures: CC BY 4.0.",
    ]
    for i, t in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t); c.font = Font(name=FONT, bold=(i == 1)); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 150

    sheets = []
    sheets.append(figure_sheet(wb, "Fig 4", "Figure 4: the construction's inputs before and after smoothing. Paper: digitised. Ours: profiles_f1.csv of the Part 1 run (the 1D construction at smoothFactor 4000). Replica: the same construction at the fitted spans 6100 / 3540.", r, [
        ("rho' before smoothing (x c^2)", P(4, "rho_unsmoothed_c2"), prof_at("rho_u"), None),
        ("rho~ smoothed (x c^2)", P(4, "rho_smoothed_c2"), prof_at("rho_s"), rep_rho),
        ("P' before smoothing", P(4, "P_unsmoothed"), prof_at("P_u"), None),
        ("P~ smoothed", P(4, "P_smoothed"), prof_at("P_s"), rep_P)], SCI))
    sheets.append(figure_sheet(wb, "Fig 5", "Figure 5: the static shell's metric along the y-axis with the Schwarzschild reference (r_s = 2GM/c^2 = %.4f m). Ours: the slice's g00 and g22 on the line; the Schwarzschild columns are the analytic -(1 - r_s/r) and 1/(1 - r_s/r), blank inside r_s." % R_S, r, [
        ("g00 shell", P(5, "g00"), st["g00"], rep_g00),
        ("g22 shell", P(5, "g22"), st["g22"], rep_g22),
        ("g00 Schwarzschild", P(5, "g00_schwarzschild"), g00_sch, None),
        ("g22 Schwarzschild", np.where(r > R_S, P(5, "g22_schwarzschild"), np.nan), g22_sch, None)], FIX))
    sheets.append(figure_sheet(wb, "Fig 6", "Figure 6: the static shell's Eulerian stress-energy along the y-axis. On this axis P_2 = P_yy is the radial pressure and P_1 = P_xx = P_3 = P_zz the tangential one. Replica: smoothed inputs, radial pressure (= P~) and the tangential pressure from the equilibrium identity, at the fitted spans.", r, [
        ("energy density rho", P(6, "rho"), st["rho"], None),
        ("rho~ smoothed input (x c^2)", P(6, "rho_smoothed_c2"), prof_at("rho_s"), rep_rho),
        ("P_2 radial", P(6, "P_2 (radial)"), st["P_yy"], rep_Pr),
        ("P_1 = P_3 tangential", P(6, "P_1 = P_3 (tangential)"), st["P_xx"], rep_Pt),
        ("P~ smoothed input", P(6, "P_smoothed"), prof_at("P_s"), rep_P)], SCI))
    sheets.append(figure_sheet(wb, "Fig 7", "Figure 7: the static shell's four sampled energy-condition maps (WarpFactory, 100 directions x 10 speeds), the paper's method. The exported Null and Weak maps lower indices with the coordinate metric (claim C9); 'consistent metric' columns are the same observer set evaluated with Minkowski, ours only.", r, [
        ("Null", P(7, "null"), st["null"], None), ("Weak", P(7, "weak"), st["weak"], None),
        ("Strong", P(7, "strong"), st["strong"], None), ("Dominant", P(7, "dominant"), st["dominant"], None),
        ("Null, consistent metric", None, st["null_fix"], None), ("Weak, consistent metric", None, st["weak_fix"], None)], SCI))
    sheets.append(figure_sheet(wb, "Fig 8", "Figure 8: the warp shell's metric along the y-axis (shift 0.02 along +x). The paper's Figure 8 is a raster image in the PDF, so there is no digitised paper column; the comparison is by eye in pairs/fig08_pair.png. Ours only.", r, [
        ("g00", None, w["g00"], None), ("g01", None, w["g01"], None), ("g22", None, w["g22"], None)], FIX))
    sheets.append(figure_sheet(wb, "Fig 9", "Figure 9: the warp shell's Eulerian stress-energy along the y-axis (shift 0.02 along +x). p_1 = p_x is the momentum density; sigma_12 = P_xy is the shear, which vanishes on the exact axis and is 1.9e37 at our line's x = -0.1 m.", r, [
        ("energy density rho", P(9, "rho"), w["rho"], None),
        ("P_2 radial", P(9, "P_2"), w["P_yy"], None), ("P_1", P(9, "P_1"), w["P_xx"], None), ("P_3", P(9, "P_3"), w["P_zz"], None),
        ("momentum p_1 (x c)", P(9, "p_1"), w["p_x"], None), ("momentum p_2 = p_3 (x c)", P(9, "p_2 = p_3"), w["p_y"], None),
        ("shear sigma_12", P(9, "sigma_12"), w["P_xy"], None)], SCI))
    sheets.append(figure_sheet(wb, "Fig 10", "Figure 10: the warp shell's four sampled energy-condition maps, as Figure 7.", r, [
        ("Null", P(10, "null"), w["null"], None), ("Weak", P(10, "weak"), w["weak"], None),
        ("Strong", P(10, "strong"), w["strong"], None), ("Dominant", P(10, "dominant"), w["dominant"], None),
        ("Null, consistent metric", None, w["null_fix"], None), ("Weak, consistent metric", None, w["weak_fix"], None)], SCI))

    # Summary: extrema per curve with formulas against the figure sheets
    ws = wb.create_sheet("Summary", 1)
    heads = ["Sheet", "Curve", "Paper: max", "Ours: max", "max ratio paper / ours", "Paper: min", "Ours: min",
             "RMS difference / max of ours", "Largest |difference| / max of ours", "Replica at the figures' smoothing: max", "Replica max / paper max"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, len(heads)); ws.row_dimensions[1].height = 45
    row = 2
    for (sheet, blocks, r0, r1) in [(s[0], s[1], s[2], s[3]) for s in sheets]:
        name = sheet.title
        for col, label, has_paper, has_rep in blocks:
            cp, co, cd, cr = (get_column_letter(col + k) for k in range(4))
            ws.cell(row=row, column=1, value=name); ws.cell(row=row, column=2, value=label)
            q = f"'{name}'!"
            if has_paper:
                ws.cell(row=row, column=3, value=f"=MAX({q}{cp}{r0}:{cp}{r1})")
                ws.cell(row=row, column=6, value=f"=MIN({q}{cp}{r0}:{cp}{r1})")
                ws.cell(row=row, column=5, value=f"=C{row}/D{row}").number_format = "0.0000"
                ws.cell(row=row, column=8, value=f"=SQRT(SUMPRODUCT(({q}{cd}{r0}:{cd}{r1})^2)/COUNT({q}{cd}{r0}:{cd}{r1}))/MAX({q}{co}{r0}:{co}{r1})").number_format = "0.0000"
                ws.cell(row=row, column=9, value=f"=MAX(MAX({q}{cd}{r0}:{cd}{r1}),-MIN({q}{cd}{r0}:{cd}{r1}))/MAX({q}{co}{r0}:{co}{r1})").number_format = "0.0000"
            ws.cell(row=row, column=4, value=f"=MAX({q}{co}{r0}:{co}{r1})")
            ws.cell(row=row, column=7, value=f"=MIN({q}{co}{r0}:{co}{r1})")
            if has_rep:
                ce = get_column_letter(col + 4)
                ws.cell(row=row, column=10, value=f"=MAX({q}{ce}{r0}:{ce}{r1})")
                if has_paper:
                    ws.cell(row=row, column=11, value=f"=J{row}/C{row}").number_format = "0.0000"
            fmt = FIX if name in ("Fig 5", "Fig 8") else SCI
            for cidx in (3, 4, 6, 7, 10):
                ws.cell(row=row, column=cidx).number_format = fmt
            row += 1
    write_note(ws, row + 1, "Ratios above 1 mean the paper's curve is larger. The 'replica at the figures' smoothing' column is the construction rebuilt at the spans fitted to the paper's inputs; where it is within a few tenths of a percent of the paper's maximum, the difference between the paper and ours is the smoothing and nothing else. Formulas reference the figure sheets; Excel evaluates them on opening.", ncols=11)
    for j in range(1, len(heads) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 18
    ws.column_dimensions["B"].width = 30
    ws.freeze_panes = "C2"
    for rr in ws.iter_rows(min_row=2, max_row=row - 1, max_col=len(heads)):
        for cell in rr:
            cell.font = Font(name=FONT)

    # Smoothing fit
    ws = wb.create_sheet("Smoothing fit")
    with open(HERE / "fit_smoothing.csv", encoding="utf-8") as f:
        rows = [l.rstrip("\n") for l in f]
    hdr = rows[0].split(","); data = [x.split(",") for x in rows[1:] if x and not x.startswith("#")]; comment = [x[2:] for x in rows if x.startswith("#")]
    for j, h in enumerate(hdr, start=1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, len(hdr)); ws.row_dimensions[1].height = 45
    for i, d in enumerate(data, start=2):
        for j, v in enumerate(d, start=1):
            try:
                val = float(v); c = ws.cell(row=i, column=j, value=val); c.number_format = SCI if abs(val) > 1e6 else "0.0000"
            except ValueError:
                ws.cell(row=i, column=j, value=v)
            ws.cell(row=i, column=j).font = Font(name=FONT)
    write_note(ws, len(data) + 3, "Spans in radial samples of the construction's grid (dr = 5.07e-4 m), four moving-average passes each. " + " ".join(comment) + " The W1 live script is the published configuration; 'fitted to Figure 6' are the spans that best reproduce the paper's digitised smoothed inputs (Figure 4 gives the same spans); 'paper, Figure 6' are the paper's digitised values of the same quantities.", ncols=9)
    for j in range(1, len(hdr) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 18

    # Part 2
    ws = wb.create_sheet("Part 2 exact vs sampled")
    keep = ["compactness", "beta", "beta_over_C", "ex_nec", "ex_wec", "ex_sec", "ex_dec", "ex_nulldef", "dense", "s_null_fix", "s_null_raw", "s_weak_raw", "s_strong", "s_dom",
            "typeI_frac", "n_ex_dec", "n_s_dom", "miss_dec_dom", "miss_nec_fix", "miss_nec_dense", "routeB_maxdiff", "r_worst_dec"]
    labels = {"compactness": "compactness 2GM/(R2 c^2)", "beta": "shift beta", "beta_over_C": "beta / C", "ex_nec": "exact NEC margin / rho0 (warpax)", "ex_wec": "exact WEC margin / rho0",
              "ex_sec": "exact SEC margin / rho0", "ex_dec": "exact DEC margin / rho0", "ex_nulldef": "exact null deficit / rho0 (min over null k)", "dense": "null, 4000 directions / rho0 (Part 1)",
              "s_null_fix": "null, 100 directions, consistent metric / rho0", "s_null_raw": "WarpFactory Null map (x2) / rho0", "s_weak_raw": "WarpFactory Weak map / rho0", "s_strong": "WarpFactory Strong map / rho0",
              "s_dom": "WarpFactory Dominant map / rho0", "typeI_frac": "Type I fraction (warpax census)", "n_ex_dec": "cells violating DEC, exact", "n_s_dom": "cells flagged by the Dominant map",
              "miss_dec_dom": "fraction of exact DEC violations the Dominant map misses", "miss_nec_fix": "fraction of exact NEC violations 100 directions miss", "miss_nec_dense": "same, 4000 directions",
              "routeB_maxdiff": "route A vs B agreement, nec/wec/sec, / rho0", "r_worst_dec": "radius of the worst DEC cell [m]"}
    with open(RUN / "exact_cases.csv", encoding="utf-8") as f:
        rd = list(csv.DictReader(f))
    for j, k in enumerate(keep, start=1):
        ws.cell(row=1, column=j, value=labels[k])
    style_header(ws, 1, len(keep)); ws.row_dimensions[1].height = 60
    for i, x in enumerate(rd, start=2):
        for j, k in enumerate(keep, start=1):
            v = float(x[k]); c = ws.cell(row=i, column=j, value=v); c.font = Font(name=FONT)
            c.number_format = "0.000E+00" if k.startswith(("ex_", "dense", "s_", "routeB")) else ("0.0000" if k in ("compactness", "beta", "beta_over_C", "typeI_frac", "miss_dec_dom", "miss_nec_fix", "miss_nec_dense") else "0.00")
    write_note(ws, len(rd) + 3, "One row per (compactness, shift) of the refined scan; margins are fractions of rho0, the peak energy density of the static shell at that compactness; a cell violates below -1e-4 rho0. Source: work/matrix/run-fine/exact_cases.csv, work/matrix/RESULTS.md section 7.", ncols=12)
    for j in range(1, len(keep) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.freeze_panes = "D2"

    out = HERE / "comparison.xlsx"
    wb.save(out)
    print("written:", out, "sheets:", wb.sheetnames)


if __name__ == "__main__":
    main()
